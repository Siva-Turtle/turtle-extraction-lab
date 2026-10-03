"""Export Attributes v2 proposal workbook (read-only DB, no migration).

Version 2 (rev 2) - decisions approved by Siva, 2026-10-03.

Run from backend/ with:
    PYTHONPATH=. uv run --with openpyxl python ../scripts/export_attributes_v2.py

Reads current attributes ordered by group_name, name and writes
exports/attributes_v2.xlsx with sheets:
  README, Attributes v2, Dict formats, Old -> New mapping, JSON schema.

Rev-2 rules:
  * 7 NEW list attributes (assets, accounts, expenses, goals,
    income_sources, insurance_policies, liabilities). Tax reverted -
    groups 'Tax' and 'Tax / Compliance' keep their v1 attributes.
  * Adequacy booleans kept as-is (Status UNCHANGED).
  * NO outer {value, confidence, confidence_type, evidence} wrapper for
    the 7 NEW list attributes - the value IS the list, per-item
    confidence only. Unchanged attributes keep the wrapper.
"""

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.db.models import Attribute, Agent, agent_attributes
from app.db.session import SessionLocal

REPO_ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = REPO_ROOT / "exports" / "attributes_v2.xlsx"

AMOUNT_DESC = (
    "Currency code or unit, one space, digits only - no commas, no symbols, no words: "
    "'INR 1000000', 'USD 5400', '50 units', '0.5 BTC'. Convert lakh/crore to digits "
    "(10 lakh -> INR 1000000). '0' when the Client says they have none of this type. "
    "'Yes' when the Client has it but gave no figure. null when not mentioned."
)

ITEM_RULE = (
    "Return one item per distinct holding/instance mentioned. If the Client mentions "
    "a type only in aggregate, return one item for that type. Return an empty list "
    "when the group is not discussed."
)

NO_WRAPPER_SENTENCE = (
    "Return the list directly - there is no outer value/confidence wrapper."
)

CONFIDENCE_DESC = "How confident the model is in this item (0-1)."
CONFIDENCE_TYPE_DESC = (
    "quoted = stated verbatim; normalized = stated but reformatted "
    "(e.g. '10 lakh' -> 'INR 1000000'); inferred = clearly implied; "
    "calculated = derived from several stated pieces; "
    "not_found = no value found for this type."
)
EVIDENCE_DESC = "Short verbatim quote(s) from the transcript supporting this item."

CONFIDENCE_ENUM = ["quoted", "normalized", "inferred", "calculated", "not_found"]

ADEQUACY_NAMES = {
    "term_insurance_coverage_adequacy",
    "health_insurance_coverage_adequacy",
}

REPLACED_GROUPS = {
    "Assets",
    "Banking / Accounts",
    "Expenses",
    "Goals",
    "Income",
    "Insurance",
    "Liabilities",
}

# ---------------------------------------------------------------------------
# New attribute definitions (proposal). Fields are in display order; the 3
# common fields are appended automatically.
# Each field: name, type (string|number|integer|enum|array), required bool,
#   nullable bool, enum list, description (short, with 1-2 inline e.g.).
# Each attribute also carries: purpose, type_field + type_gloss (meanings for
# non-obvious enum values), and 3 transcript examples (transcript + items).
# The full self-contained "Description / Definition" is built by
# build_long_description() below.
# ---------------------------------------------------------------------------

COMMON_FIELDS = [
    {
        "name": "confidence",
        "type": "number",
        "required": True,
        "nullable": False,
        "enum": [],
        "description": CONFIDENCE_DESC,
    },
    {
        "name": "confidence_type",
        "type": "enum",
        "required": True,
        "nullable": False,
        "enum": list(CONFIDENCE_ENUM),
        "description": CONFIDENCE_TYPE_DESC,
    },
    {
        "name": "evidence",
        "type": "string",
        "required": False,
        "nullable": True,
        "enum": [],
        "description": EVIDENCE_DESC,
    },
]


def _f(name, ftype, required, nullable, enum=None, description=""):
    return {
        "name": name,
        "type": ftype,
        "required": required,
        "nullable": nullable,
        "enum": list(enum or []),
        "description": description,
    }


NEW_ATTRIBUTES = [
    {
        "name": "assets",
        "group": "Assets",
        "agent": "asset",
        "purpose": (
            "This attribute consolidates every asset holding the Client mentions "
            "into a single list so the planner sees the full portfolio in one place."
        ),
        "type_field": "asset_type",
        "type_gloss": (
            "Meanings: unlisted_stocks = shares not listed on an exchange "
            "(e.g. ESOPs, startup equity); deposit = fixed/recurring deposits held "
            "outside bank accounts listed in accounts; personal_debt = money the "
            "Client has lent to someone; others = any type not listed; total = an "
            "explicitly stated total for the group."
        ),
        "fields": [
            _f("asset_type", "enum", True, False, [
                "stocks", "unlisted_stocks", "equity_mutual_fund", "equity_etf",
                "commodity_etf", "debt_etf", "deposit", "cash", "debt_mf",
                "personal_debt", "bonds", "reits", "real_estate", "crypto",
                "others", "total",
            ], "Type of asset holding, e.g. 'stocks', 'bonds', 'real_estate'."),
            _f("name", "string", False, True, [],
               "Instrument / scheme / property / borrower name, e.g. 'HDFC Flexi Cap', 'Flat in Whitefield, Bengaluru'."),
            _f("location", "enum", False, True, ["India", "Foreign"],
               "Where the asset is held, e.g. 'India'. For asset_type total, null = overall total across locations."),
            _f("amount", "string", False, True, [],
               "Current value, or units held, e.g. 'INR 500000'. " + AMOUNT_DESC),
        ],
        "examples": [
            {
                "transcript": "I have about 5 lakh in RBI bonds and around 2 lakh in HDFC Flexi Cap.",
                "items": [
                    {
                        "asset_type": "bonds",
                        "name": "RBI bonds",
                        "location": "India",
                        "amount": "INR 500000",
                        "confidence": 0.9,
                        "confidence_type": "normalized",
                        "evidence": "I have about 5 lakh in RBI bonds",
                    },
                    {
                        "asset_type": "equity_mutual_fund",
                        "name": "HDFC Flexi Cap",
                        "location": "India",
                        "amount": "INR 200000",
                        "confidence": 0.9,
                        "confidence_type": "normalized",
                        "evidence": "around 2 lakh in HDFC Flexi Cap",
                    },
                ],
            },
            {
                "transcript": "I hold about $5,400 of Apple shares in my US brokerage.",
                "items": [
                    {
                        "asset_type": "stocks",
                        "name": "Apple shares",
                        "location": "Foreign",
                        "amount": "USD 5400",
                        "confidence": 0.9,
                        "confidence_type": "normalized",
                        "evidence": "I hold about $5,400 of Apple shares in my US brokerage",
                    },
                ],
            },
            {
                "transcript": "I do hold some crypto but I don't track the value, and I own no real estate.",
                "items": [
                    {
                        "asset_type": "crypto",
                        "name": None,
                        "location": None,
                        "amount": "Yes",
                        "confidence": 0.8,
                        "confidence_type": "quoted",
                        "evidence": "I do hold some crypto but I don't track the value",
                    },
                    {
                        "asset_type": "real_estate",
                        "name": None,
                        "location": None,
                        "amount": "0",
                        "confidence": 0.95,
                        "confidence_type": "quoted",
                        "evidence": "I own no real estate",
                    },
                ],
            },
        ],
    },
    {
        "name": "accounts",
        "group": "Banking / Accounts",
        "agent": "account",
        "purpose": (
            "This attribute consolidates every account the Client mentions - bank, "
            "investment, government schemes, brokers and custody - into a single "
            "list so balances can be totalled per location."
        ),
        "type_field": "account_type",
        "type_gloss": (
            "Meanings: ppf = Public Provident Fund; bank_deposit = fixed/recurring "
            "deposit held at a bank; nps = National Pension System; ssy = Sukanya "
            "Samriddhi Yojana; epf = Employees' Provident Fund; pms = Portfolio "
            "Management Service; aif = Alternative Investment Fund; crypto_broker = "
            "exchange or broker account for crypto; self_custody = crypto held in "
            "the Client's own wallet; others = any type not listed."
        ),
        "fields": [
            _f("account_type", "enum", True, False, [
                "investment_account", "bank_account", "ppf", "bank_deposit",
                "nps", "ssy", "epf", "pms", "aif", "crypto_broker",
                "self_custody", "others",
            ], "Type of account, e.g. 'bank_account', 'ppf'."),
            _f("provider", "string", False, True, [],
               "Bank / broker / fund house / wallet name, e.g. 'HDFC Bank', 'Zerodha'."),
            _f("account_name", "string", False, True, [],
               "Specific account or product name if given, e.g. 'Salary account'."),
            _f("location", "enum", False, True, ["India", "Foreign"],
               "Where the account is held, e.g. 'India'."),
            _f("balance", "string", False, True, [],
               "Account balance, e.g. 'INR 250000'. " + AMOUNT_DESC),
        ],
        "examples": [
            {
                "transcript": "I have around 2.5 lakh in my HDFC salary account and about 8 lakh in PPF.",
                "items": [
                    {
                        "account_type": "bank_account",
                        "provider": "HDFC Bank",
                        "account_name": "Salary account",
                        "location": "India",
                        "balance": "INR 250000",
                        "confidence": 0.9,
                        "confidence_type": "normalized",
                        "evidence": "around 2.5 lakh in my HDFC salary account",
                    },
                    {
                        "account_type": "ppf",
                        "provider": None,
                        "account_name": None,
                        "location": "India",
                        "balance": "INR 800000",
                        "confidence": 0.9,
                        "confidence_type": "normalized",
                        "evidence": "about 8 lakh in PPF",
                    },
                ],
            },
            {
                "transcript": "I keep about ten thousand dollars in my Chase savings account in the US.",
                "items": [
                    {
                        "account_type": "bank_account",
                        "provider": "Chase",
                        "account_name": "Savings account",
                        "location": "Foreign",
                        "balance": "USD 10000",
                        "confidence": 0.9,
                        "confidence_type": "normalized",
                        "evidence": "about ten thousand dollars in my Chase savings account in the US",
                    },
                ],
            },
            {
                "transcript": "I do have an NPS account but I don't remember the balance, and I have no crypto account.",
                "items": [
                    {
                        "account_type": "nps",
                        "provider": None,
                        "account_name": None,
                        "location": "India",
                        "balance": "Yes",
                        "confidence": 0.8,
                        "confidence_type": "quoted",
                        "evidence": "I do have an NPS account but I don't remember the balance",
                    },
                    {
                        "account_type": "crypto_broker",
                        "provider": None,
                        "account_name": None,
                        "location": None,
                        "balance": "0",
                        "confidence": 0.95,
                        "confidence_type": "quoted",
                        "evidence": "I have no crypto account",
                    },
                ],
            },
        ],
    },
    {
        "name": "expenses",
        "group": "Expenses",
        "agent": "expense",
        "purpose": (
            "This attribute consolidates every expense the Client mentions, "
            "recurring or one-time, into a single list so the planner sees the "
            "full outflow picture."
        ),
        "type_field": "expense_type",
        "type_gloss": (
            "Meanings: loan_emis_* = EMI paid for that loan type; loan_emis_total = "
            "an explicitly stated total across EMIs; charity_donations = giving; "
            "miscellaneous = anything not listed; total = an explicitly stated "
            "total for the group."
        ),
        "fields": [
            _f("expense_type", "enum", True, False, [
                "rent", "utilities", "groceries", "household_staff",
                "childcare_school_fees", "entertainment_dining",
                "transport_fuel", "family_support", "shopping_purchases",
                "subscriptions_memberships", "insurance_premiums",
                "travel_vacations", "medical_expenses", "loan_emis_education",
                "loan_emis_personal", "loan_emis_home", "loan_emis_vehicle",
                "loan_emis_others", "loan_emis_total", "charity_donations",
                "miscellaneous", "total",
            ], "Type of expense, e.g. 'rent', 'travel_vacations'."),
            _f("description", "string", False, True, [],
               "What the expense is for, e.g. '2BHK in Whitefield, Bengaluru'."),
            _f("amount", "string", False, True, [],
               "Amount per period, e.g. 'INR 30000'. " + AMOUNT_DESC),
            _f("frequency", "enum", False, True, ["monthly", "annual", "one_time"],
               "How often it recurs, e.g. 'monthly'. Keep the Client's own period - do NOT convert annual to monthly."),
        ],
        "examples": [
            {
                "transcript": "We pay 30k rent for our 2BHK in Whitefield and 5k for utilities every month.",
                "items": [
                    {
                        "expense_type": "rent",
                        "description": "2BHK in Whitefield, Bengaluru",
                        "amount": "INR 30000",
                        "frequency": "monthly",
                        "confidence": 0.9,
                        "confidence_type": "normalized",
                        "evidence": "We pay 30k rent for our 2BHK in Whitefield",
                    },
                    {
                        "expense_type": "utilities",
                        "description": None,
                        "amount": "INR 5000",
                        "frequency": "monthly",
                        "confidence": 0.9,
                        "confidence_type": "normalized",
                        "evidence": "5k for utilities every month",
                    },
                ],
            },
            {
                "transcript": "We spend about 2 lakh a year on travel and vacations.",
                "items": [
                    {
                        "expense_type": "travel_vacations",
                        "description": None,
                        "amount": "INR 200000",
                        "frequency": "annual",
                        "confidence": 0.9,
                        "confidence_type": "normalized",
                        "evidence": "about 2 lakh a year on travel and vacations",
                    },
                ],
            },
            {
                "transcript": "Groceries I don't track exactly, and we pay zero rent since we own the house.",
                "items": [
                    {
                        "expense_type": "groceries",
                        "description": None,
                        "amount": "Yes",
                        "frequency": None,
                        "confidence": 0.8,
                        "confidence_type": "quoted",
                        "evidence": "Groceries I don't track exactly",
                    },
                    {
                        "expense_type": "rent",
                        "description": "Own house, no rent",
                        "amount": "0",
                        "frequency": "monthly",
                        "confidence": 0.9,
                        "confidence_type": "inferred",
                        "evidence": "we pay zero rent since we own the house",
                    },
                ],
            },
        ],
    },
    {
        "name": "goals",
        "group": "Goals",
        "agent": "goal",
        "purpose": (
            "This attribute consolidates every goal the Client mentions into a "
            "single list so the planner can size, time and prioritise each one."
        ),
        "type_field": "goal_type",
        "type_gloss": (
            "Meanings: planning_a_child = expecting- or first-child costs; "
            "child_s_college_education / child_s_higher_education = UG / PG or "
            "study-abroad costs; buying_a_house_india = India home purchase; "
            "financial_freedom_fire = retire-early corpus; legacy_inheritance = "
            "leave-behind for heirs; other_goals = any goal not listed."
        ),
        "fields": [
            _f("goal_type", "enum", True, False, [
                "planning_a_child", "child_s_college_education",
                "child_s_higher_education", "buying_a_house_india",
                "buying_a_vehicle", "wedding_expenses",
                "parents_healthcare_fund", "vacation_travel_fund",
                "legacy_inheritance", "financial_freedom_fire", "other_goals",
            ], "Type of goal, e.g. 'buying_a_house_india', 'financial_freedom_fire'."),
            _f("description", "string", False, True, [],
               "Goal in the Client's own terms, e.g. \"Daughter's MS in the US\"."),
            _f("target_amount", "string", False, True, [],
               "Target corpus / cost, e.g. 'INR 8000000'. " + AMOUNT_DESC),
            _f("target_year", "integer", False, True, [],
               "Calendar year the Client wants to achieve the goal, e.g. 2032. "
               "If stated as 'in N years', add N to the meeting year and use "
               "confidence_type calculated."),
            _f("priority", "enum", False, True, ["high", "medium", "low"],
               "How important the goal is, e.g. 'high'. Only when the Client indicates it."),
        ],
        "examples": [
            {
                "transcript": "We want around 80 lakh for our daughter's MS in the US by 2032, and about 1 crore for a house in India.",
                "items": [
                    {
                        "goal_type": "child_s_higher_education",
                        "description": "Daughter's MS in the US",
                        "target_amount": "INR 8000000",
                        "target_year": 2032,
                        "priority": None,
                        "confidence": 0.85,
                        "confidence_type": "normalized",
                        "evidence": "around 80 lakh for our daughter's MS in the US by 2032",
                    },
                    {
                        "goal_type": "buying_a_house_india",
                        "description": "House in India",
                        "target_amount": "INR 10000000",
                        "target_year": None,
                        "priority": None,
                        "confidence": 0.85,
                        "confidence_type": "normalized",
                        "evidence": "about 1 crore for a house in India",
                    },
                ],
            },
            {
                "transcript": "We want to retire in 15 years with a 5 crore corpus. (Meeting year 2026.)",
                "items": [
                    {
                        "goal_type": "financial_freedom_fire",
                        "description": "Retire with a 5 crore corpus",
                        "target_amount": "INR 50000000",
                        "target_year": 2041,
                        "priority": None,
                        "confidence": 0.8,
                        "confidence_type": "calculated",
                        "evidence": "retire in 15 years with a 5 crore corpus",
                    },
                ],
            },
            {
                "transcript": "We are planning for a child but haven't fixed any amount yet.",
                "items": [
                    {
                        "goal_type": "planning_a_child",
                        "description": "Planning for a child",
                        "target_amount": "Yes",
                        "target_year": None,
                        "priority": None,
                        "confidence": 0.85,
                        "confidence_type": "quoted",
                        "evidence": "planning for a child but haven't fixed any amount yet",
                    },
                ],
            },
        ],
    },
    {
        "name": "income_sources",
        "group": "Income",
        "agent": "income",
        "purpose": (
            "This attribute consolidates every income source the Client mentions "
            "into a single list so total inflow can be compared against expenses."
        ),
        "type_field": "income_type",
        "type_gloss": (
            "Meanings: capital_gains = profit from selling assets; rental = rent "
            "received; speculation = F&O, intraday or lottery-type gains; "
            "other_income = any source not listed."
        ),
        "fields": [
            _f("income_type", "enum", True, False, [
                "salary", "business", "capital_gains", "rental",
                "speculation", "other_income",
            ], "Type of income, e.g. 'salary', 'rental'."),
            _f("source", "string", False, True, [],
               "Employer / business / property name, e.g. 'Infosys', 'Pune flat rent'."),
            _f("amount", "string", False, True, [],
               "Amount per period, e.g. 'INR 200000'. " + AMOUNT_DESC),
            _f("frequency", "enum", False, True, ["monthly", "annual", "one_time"],
               "Pay frequency, e.g. 'monthly'."),
            _f("location", "enum", False, True, ["India", "Foreign"],
               "Where the income is earned, e.g. 'India'."),
        ],
        "examples": [
            {
                "transcript": "My monthly salary at Infosys is 2 lakh and we get 25k rent from our Pune flat.",
                "items": [
                    {
                        "income_type": "salary",
                        "source": "Infosys",
                        "amount": "INR 200000",
                        "frequency": "monthly",
                        "location": "India",
                        "confidence": 0.95,
                        "confidence_type": "normalized",
                        "evidence": "My monthly salary at Infosys is 2 lakh",
                    },
                    {
                        "income_type": "rental",
                        "source": "Pune flat",
                        "amount": "INR 25000",
                        "frequency": "monthly",
                        "location": "India",
                        "confidence": 0.9,
                        "confidence_type": "normalized",
                        "evidence": "we get 25k rent from our Pune flat",
                    },
                ],
            },
            {
                "transcript": "I earn about $5,400 a month from my US contracting work.",
                "items": [
                    {
                        "income_type": "salary",
                        "source": "US contracting work",
                        "amount": "USD 5400",
                        "frequency": "monthly",
                        "location": "Foreign",
                        "confidence": 0.9,
                        "confidence_type": "normalized",
                        "evidence": "about $5,400 a month from my US contracting work",
                    },
                ],
            },
            {
                "transcript": "My wife has some business income but I don't know the figure, and I had no capital gains this year.",
                "items": [
                    {
                        "income_type": "business",
                        "source": "Wife's business",
                        "amount": "Yes",
                        "frequency": None,
                        "location": "India",
                        "confidence": 0.8,
                        "confidence_type": "quoted",
                        "evidence": "My wife has some business income but I don't know the figure",
                    },
                    {
                        "income_type": "capital_gains",
                        "source": None,
                        "amount": "0",
                        "frequency": "annual",
                        "location": None,
                        "confidence": 0.9,
                        "confidence_type": "quoted",
                        "evidence": "I had no capital gains this year",
                    },
                ],
            },
        ],
    },
    {
        "name": "insurance_policies",
        "group": "Insurance",
        "agent": "tax_and_insurance",
        "purpose": (
            "This attribute consolidates every insurance policy the Client mentions "
            "into a single list so cover gaps can be assessed. The two boolean "
            "adequacy judgements (term_insurance_coverage_adequacy, "
            "health_insurance_coverage_adequacy) stay as separate attributes."
        ),
        "type_field": "insurance_type",
        "type_gloss": (
            "Meanings: life_term = pure protection term cover; health = medical "
            "cover; ulip = Unit Linked Insurance Plan (investment-cum-insurance); "
            "other = any policy not listed."
        ),
        "fields": [
            _f("insurance_type", "enum", True, False,
               ["life_term", "health", "ulip", "other"],
               "Type of policy, e.g. 'life_term', 'health'."),
            _f("insurer", "string", False, True, [],
               "Insurance company, e.g. 'HDFC Life', 'Star Health'."),
            _f("policy_name", "string", False, True, [],
               "Plan name if given, e.g. 'Click 2 Protect'."),
            _f("covered_members", "string", False, True, [],
               "Who is covered, e.g. 'self', 'self, spouse, 2 children'."),
            _f("cover_amount", "string", False, True, [],
               "Sum assured / cover, e.g. 'INR 10000000'. " + AMOUNT_DESC),
            _f("premium_amount", "string", False, True, [],
               "Premium per period, e.g. 'INR 12000'. " + AMOUNT_DESC),
            _f("premium_frequency", "enum", False, True,
               ["monthly", "quarterly", "annual", "one_time"],
               "How often the premium is paid, e.g. 'annual'."),
            _f("source", "enum", False, True, ["employer", "personal"],
               "Who provides the cover, e.g. 'personal'. employer = group / corporate cover."),
        ],
        "examples": [
            {
                "transcript": "I have a 1 crore HDFC term plan with 12k annual premium, and a 10 lakh Star Health family floater.",
                "items": [
                    {
                        "insurance_type": "life_term",
                        "insurer": "HDFC Life",
                        "policy_name": None,
                        "covered_members": "self",
                        "cover_amount": "INR 10000000",
                        "premium_amount": "INR 12000",
                        "premium_frequency": "annual",
                        "source": "personal",
                        "confidence": 0.9,
                        "confidence_type": "normalized",
                        "evidence": "1 crore HDFC term plan with 12k annual premium",
                    },
                    {
                        "insurance_type": "health",
                        "insurer": "Star Health",
                        "policy_name": None,
                        "covered_members": "family",
                        "cover_amount": "INR 1000000",
                        "premium_amount": None,
                        "premium_frequency": None,
                        "source": "personal",
                        "confidence": 0.85,
                        "confidence_type": "normalized",
                        "evidence": "a 10 lakh Star Health family floater",
                    },
                ],
            },
            {
                "transcript": "My employer gives me a 5 lakh health cover for my family; I pay nothing for it.",
                "items": [
                    {
                        "insurance_type": "health",
                        "insurer": None,
                        "policy_name": None,
                        "covered_members": "family",
                        "cover_amount": "INR 500000",
                        "premium_amount": "0",
                        "premium_frequency": None,
                        "source": "employer",
                        "confidence": 0.9,
                        "confidence_type": "quoted",
                        "evidence": "My employer gives me a 5 lakh health cover for my family",
                    },
                ],
            },
            {
                "transcript": "I do have a ULIP from years ago but I don't remember the cover, and I have no separate term cover of my own.",
                "items": [
                    {
                        "insurance_type": "ulip",
                        "insurer": None,
                        "policy_name": None,
                        "covered_members": None,
                        "cover_amount": "Yes",
                        "premium_amount": None,
                        "premium_frequency": None,
                        "source": "personal",
                        "confidence": 0.75,
                        "confidence_type": "quoted",
                        "evidence": "I do have a ULIP from years ago but I don't remember the cover",
                    },
                    {
                        "insurance_type": "life_term",
                        "insurer": None,
                        "policy_name": None,
                        "covered_members": None,
                        "cover_amount": "0",
                        "premium_amount": None,
                        "premium_frequency": None,
                        "source": None,
                        "confidence": 0.9,
                        "confidence_type": "quoted",
                        "evidence": "I have no separate term cover of my own",
                    },
                ],
            },
        ],
    },
    {
        "name": "liabilities",
        "group": "Liabilities",
        "agent": "liability",
        "purpose": (
            "This attribute consolidates every liability or loan the Client mentions "
            "into a single list so outstanding debt and EMI load can be totalled."
        ),
        "type_field": "liability_type",
        "type_gloss": (
            "Meanings: credit_card_debt = unpaid card dues; other_loan = any loan "
            "not listed; total = an explicitly stated total for the group."
        ),
        "fields": [
            _f("liability_type", "enum", True, False, [
                "home_loan", "vehicle_loan", "education_loan",
                "personal_loan", "gold_loan", "credit_card_debt",
                "other_loan", "total",
            ], "Type of loan, e.g. 'home_loan', 'credit_card_debt'."),
            _f("lender", "string", False, True, [],
               "Lender name, e.g. 'SBI', 'HDFC Bank'."),
            _f("outstanding_amount", "string", False, True, [],
               "Amount still owed, e.g. 'INR 4500000'. " + AMOUNT_DESC),
            _f("emi_amount", "string", False, True, [],
               "Monthly EMI, e.g. 'INR 42000'. " + AMOUNT_DESC),
            _f("interest_rate", "string", False, True, [],
               "Rate as stated, e.g. '8.5%'."),
            _f("end_year", "integer", False, True, [],
               "Year the loan ends, e.g. 2040."),
        ],
        "examples": [
            {
                "transcript": "I have an SBI home loan of 45 lakh outstanding, EMI of 42 thousand at 8.5%, till 2040.",
                "items": [
                    {
                        "liability_type": "home_loan",
                        "lender": "SBI",
                        "outstanding_amount": "INR 4500000",
                        "emi_amount": "INR 42000",
                        "interest_rate": "8.5%",
                        "end_year": 2040,
                        "confidence": 0.9,
                        "confidence_type": "normalized",
                        "evidence": "SBI home loan of 45 lakh outstanding, EMI of 42 thousand at 8.5%, till 2040",
                    },
                ],
            },
            {
                "transcript": "I have a 3 lakh personal loan and about 50 thousand outstanding on my credit card.",
                "items": [
                    {
                        "liability_type": "personal_loan",
                        "lender": None,
                        "outstanding_amount": "INR 300000",
                        "emi_amount": None,
                        "interest_rate": None,
                        "end_year": None,
                        "confidence": 0.9,
                        "confidence_type": "normalized",
                        "evidence": "I have a 3 lakh personal loan",
                    },
                    {
                        "liability_type": "credit_card_debt",
                        "lender": None,
                        "outstanding_amount": "INR 50000",
                        "emi_amount": None,
                        "interest_rate": None,
                        "end_year": None,
                        "confidence": 0.85,
                        "confidence_type": "normalized",
                        "evidence": "about 50 thousand outstanding on my credit card",
                    },
                ],
            },
            {
                "transcript": "I still have some education loan left but I'm not sure how much, and I have no vehicle loan.",
                "items": [
                    {
                        "liability_type": "education_loan",
                        "lender": None,
                        "outstanding_amount": "Yes",
                        "emi_amount": None,
                        "interest_rate": None,
                        "end_year": None,
                        "confidence": 0.8,
                        "confidence_type": "quoted",
                        "evidence": "I still have some education loan left but I'm not sure how much",
                    },
                    {
                        "liability_type": "vehicle_loan",
                        "lender": None,
                        "outstanding_amount": "0",
                        "emi_amount": None,
                        "interest_rate": None,
                        "end_year": None,
                        "confidence": 0.95,
                        "confidence_type": "quoted",
                        "evidence": "I have no vehicle loan",
                    },
                ],
            },
        ],
    },
]


def build_long_description(purpose, fields, type_field, type_gloss, examples):
    """Self-contained definition: purpose + item rule + no-wrapper sentence,
    then one Properties line per item property, then transcript examples."""
    para = purpose.strip() + " " + ITEM_RULE + " " + NO_WRAPPER_SENTENCE
    prop_lines = []
    for f in fields:
        req = "required" if f.get("required") else "nullable"
        line = "- {} ({}, {}): {}".format(
            f["name"], f["type"], req, (f.get("description") or "").strip()
        )
        enum = [v for v in (f.get("enum", []) or []) if v is not None]
        if enum:
            line += " Allowed values: {}.".format(", ".join(str(v) for v in enum))
        if f["name"] == type_field and type_gloss:
            line += " " + type_gloss.strip()
        prop_lines.append(line)
    ex_lines = []
    for i, ex in enumerate(examples, start=1):
        js = json.dumps(ex["items"], indent=2, ensure_ascii=False)
        ex_lines.append('Example {} - Client: "{}" -> {}'.format(i, ex["transcript"], js))
    return para + "\n\nProperties:\n" + "\n".join(prop_lines) + "\n\nExamples:\n" + "\n".join(ex_lines)


for _na in NEW_ATTRIBUTES:
    _na["fields"] = list(_na["fields"]) + [dict(c) for c in COMMON_FIELDS]
    _na["description"] = build_long_description(
        _na["purpose"], _na["fields"], _na["type_field"], _na["type_gloss"], _na["examples"]
    )

NEW_BY_NAME = {n["name"]: n for n in NEW_ATTRIBUTES}

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
NEW_FILL = PatternFill("solid", fgColor="C6EFCE")
BLOCK_FILL = PatternFill("solid", fgColor="DDEBF7")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER_TOP_WRAP = Alignment(wrap_text=True, vertical="top")


def clean_cell(v):
    """Blank cells instead of the string 'None' / None values."""
    if v is None:
        return ""
    if isinstance(v, str) and v.strip() == "None":
        return ""
    return v


def style_header_row(ws, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_TOP_WRAP
    ws.freeze_panes = "A2"


def set_widths(ws, widths):
    from openpyxl.utils import get_column_letter

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def format_enum_values(values):
    if not values:
        return ""
    cleaned = [v for v in values if v is not None and str(v).strip() != "" and str(v) != "None"]
    if not cleaned:
        return ""
    return " | ".join(str(v) for v in cleaned)


def format_db_item_fields(attr):
    """One per line as 'name: type [enum] (nullable/required)' for DB rows."""
    props = []
    atype = (attr.type or "").strip().lower()
    if atype == "object":
        props = attr.object_properties or []
    elif atype == "array":
        raw = attr.array_items or {}
        if isinstance(raw, dict) and str(raw.get("kind", "")).lower() == "object":
            props = raw.get("properties", []) or []
        else:
            return ""
    else:
        return ""
    lines = []
    for p in props:
        if not isinstance(p, dict):
            continue
        name = p.get("name")
        if name is None or str(name) == "None":
            continue
        name = str(name)
        ptype = p.get("type", "string")
        if ptype is None or str(ptype) == "None":
            ptype = "string"
        ptype = str(ptype)
        enum = [v for v in (p.get("enum", []) or []) if v is not None and str(v) != "None"]
        null_allowed = p.get("null_allowed", True)
        suffix = "(nullable)" if null_allowed else "(required)"
        if enum:
            lines.append("{}: {} [{}] {}".format(name, ptype, " | ".join(str(v) for v in enum), suffix))
        else:
            lines.append("{}: {} {}".format(name, ptype, suffix))
    return "\n".join(lines)


def format_proposal_item_fields(fields):
    lines = []
    for f in fields:
        enum = [v for v in (f.get("enum", []) or []) if v is not None and str(v) != "None"]
        req = "(required)" if f.get("required") else "(nullable)"
        if f["type"] == "enum" and enum:
            lines.append("{}: enum [{}] {}".format(f['name'], " | ".join(str(v) for v in enum), req))
        elif enum:
            lines.append("{}: {} [{}] {}".format(f['name'], f['type'], " | ".join(str(v) for v in enum), req))
        else:
            lines.append("{}: {} {}".format(f['name'], f['type'], req))
    return "\n".join(lines)


def field_to_json_schema(f):
    desc = f.get("description", "")
    ftype = f.get("type", "string")
    nullable = bool(f.get("nullable", True))
    enum = [v for v in (f.get("enum", []) or []) if v is not None and str(v) != "None"]
    if ftype == "enum":
        if nullable:
            return {
                "type": ["string", "null"],
                "enum": list(enum) + [None],
                "description": desc,
            }
        return {"type": "string", "enum": list(enum), "description": desc}
    if ftype == "array":
        if nullable:
            return {
                "type": ["array", "null"],
                "items": {"type": "string"},
                "description": desc,
            }
        return {
            "type": "array",
            "items": {"type": "string"},
            "description": desc,
        }
    if ftype == "integer":
        if nullable:
            return {"type": ["integer", "null"], "description": desc}
        return {"type": "integer", "description": desc}
    if ftype == "number":
        if nullable:
            return {"type": ["number", "null"], "description": desc}
        return {"type": "number", "description": desc}
    # string (covers amount/balance/evidence/... stored as strings)
    if nullable:
        return {"type": ["string", "null"], "description": desc}
    return {"type": "string", "description": desc}


def build_json_schema(new_attr):
    props = {}
    required = []
    for f in new_attr["fields"]:
        props[f["name"]] = field_to_json_schema(f)
        required.append(f["name"])
    return {
        "type": "array",
        "items": {
            "type": "object",
            "properties": props,
            "required": required,
            "additionalProperties": False,
        },
    }


def mapping_for_old(group, name):
    if group == "Assets":
        if name == "others_asset_type":
            return ("assets", "asset_type=others")
        return ("assets", "asset_type={}".format(name))
    if group == "Banking / Accounts":
        if name == "others_account_type":
            return ("accounts", "account_type=others")
        return ("accounts", "account_type={}".format(name))
    if group == "Expenses":
        return ("expenses", "expense_type={}".format(name))
    if group == "Goals":
        return ("goals", "goal_type={}".format(name))
    if group == "Income":
        return ("income_sources", "income_type={}".format(name))
    if group == "Insurance":
        if name == "health":
            return ("insurance_policies", "insurance_type=health")
        if name == "life_term":
            return ("insurance_policies", "insurance_type=life_term")
        if name == "other_ulip":
            return ("insurance_policies", "insurance_type=ulip or other")
    if group == "Liabilities":
        if name == "total_liabilities":
            return ("liabilities", "liability_type=total")
        return ("liabilities", "liability_type={}".format(name))
    return ("", "")


def main():
    db = SessionLocal()
    try:
        attrs = (
            db.query(Attribute)
            .order_by(Attribute.group_name, Attribute.name)
            .all()
        )
        agent_map = {a.id: a.name for a in db.query(Agent).all()}
        attr_agents = {}
        for r in attrs:
            rows = db.execute(
                agent_attributes.select().where(
                    agent_attributes.c.attribute_id == r.id
                )
            ).all()
            attr_agents[r.id] = sorted(
                {agent_map.get(x.agent_id, x.agent_id) for x in rows}
            )
    finally:
        db.close()

    total_now = len(attrs)
    deleted_rows = [
        r for r in attrs
        if r.group_name in REPLACED_GROUPS and r.name not in ADEQUACY_NAMES
    ]
    n_deleted = len(deleted_rows)
    n_after = total_now - n_deleted + len(NEW_ATTRIBUTES)

    wb = Workbook()

    # ---------------- README ----------------
    ws = wb.active
    ws.title = "README"
    ws.append(["Section", "Text"])
    style_header_row(ws, 2)
    set_widths(ws, [34, 130])
    readme_rows = [
        ("About this file",
         "Attributes v2 - a PROPOSAL for Siva to approve BEFORE any DB migration. "
         "Version 2 (rev 2). "
         "Nothing here is in the DB yet; the 7 NEW list attributes are proposed additions."),
        ("What changes",
         "The attributes of 7 groups (Assets, Banking / Accounts, Expenses, Goals, "
         "Income, Insurance, Liabilities) are replaced by ONE "
         "list attribute per group (7 new attributes). Every other attribute stays "
         "exactly as it is in the DB - including groups 'Tax' and 'Tax / Compliance' "
         "(Tax reverted, v1 attributes kept separate), group 'Banking' (credit_cards, "
         "basic_info agent) and every other group."),
        ("New attributes (7)",
         "assets (Assets, agent asset); accounts (Banking / Accounts, agent account); "
         "expenses (Expenses, agent expense); goals (Goals, agent goal); "
         "income_sources (Income, agent income); insurance_policies (Insurance, agent "
         "tax_and_insurance); liabilities (Liabilities, agent liability). "
         "All new attributes: type array, array_items kind object."),
        ("Result contract (no outer wrapper for NEW lists)",
         "For the 7 NEW list attributes the result contract wrapper "
         "{value, confidence, confidence_type, evidence} is NOT used - the attribute "
         "value IS the list, and confidence / confidence_type / evidence live only "
         "on each item. All UNCHANGED attributes keep the wrapper."),
        ("Amount convention",
         "Every field whose name ends in amount or is called balance: type string, "
         "nullable. " + AMOUNT_DESC),
        ("Common fields (per ITEM, added LAST to every item)",
         "confidence: number 0-1, required. " + CONFIDENCE_DESC + " | "
         "confidence_type: enum [quoted | normalized | inferred | calculated | not_found], "
         "required. " + CONFIDENCE_TYPE_DESC + " | "
         "evidence: string, nullable. " + EVIDENCE_DESC),
        ("Item-level rule",
         "Appended to each new attribute's definition: " + ITEM_RULE + " " + NO_WRAPPER_SENTENCE),
        ("Decisions (approved by Siva, 2026-10-03)",
         "(1) Insurance adequacy booleans (term_insurance_coverage_adequacy, "
         "health_insurance_coverage_adequacy) are KEPT as-is. "
         "(2) All proposed extra item fields (name, provider, account_name, location, "
         "description, frequency, priority, insurer, policy_name, covered_members, "
         "premium_*, source, lender, emi_amount, interest_rate, end_year) are ACCEPTED. "
         "(3) Tax REVERTED - groups 'Tax' and 'Tax / Compliance' keep their current "
         "(v1) separate attributes; there is no tax_items attribute. "
         "(4) NO outer confidence wrapper for the 7 NEW list attributes - the value "
         "IS the list, confidence lives only on items; unchanged attributes keep "
         "the wrapper. "
         "(5) confidence_type uses all five values (quoted, normalized, inferred, "
         "calculated, not_found) - note the current extraction schema enum is "
         "missing 'calculated'."),
        ("Counts",
         "Attributes in DB now: {}. Attributes deleted by this proposal: {} "
         "(every attribute in the 7 groups except the 2 adequacy booleans, which are kept). "
         "Attributes after (proposed catalogue): {} "
         "({} unchanged incl. 2 kept adequacy booleans + {} NEW).".format(
             total_now, n_deleted, n_after,
             total_now - n_deleted, len(NEW_ATTRIBUTES))),
        ("Sheets",
         "Attributes v2 = full proposed catalogue sorted by Group then Name. "
         "Dict formats = one block per new attribute with field table + 3 JSON examples. "
         "Old -> New mapping = one row per deleted attribute + 2 kept-as-is rows. "
         "JSON schema = strict OpenRouter json_schema fragment per new attribute."),
    ]
    for section, text in readme_rows:
        ws.append([section, text])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=2):
        row[0].alignment = WRAP_TOP
        row[1].alignment = WRAP_TOP
    ws.auto_filter.ref = "A1:B1"
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------------- Attributes v2 ----------------
    ws2 = wb.create_sheet("Attributes v2")
    headers2 = ["Status", "Group", "Name", "Type", "Description / Definition",
                "Enum values", "Object / array item fields", "Agents", "ID"]
    ws2.append(headers2)
    style_header_row(ws2, len(headers2))
    set_widths(ws2, [14, 22, 32, 12, 90, 30, 62, 20, 38])

    # Build combined sorted rows
    combined = []
    for r in attrs:
        if r.group_name in REPLACED_GROUPS and r.name not in ADEQUACY_NAMES:
            continue  # deleted, skip (except adequacy booleans, which are kept)
        combined.append({
            "status": "UNCHANGED",
            "group": r.group_name or "",
            "name": r.name or "",
            "type": r.type or "",
            "desc": r.description or "",
            "enum": format_enum_values(r.enum_values or []),
            "items": format_db_item_fields(r),
            "agents": ", ".join(attr_agents.get(r.id, [])),
            "id": r.id,
            "is_new": False,
        })
    for n in NEW_ATTRIBUTES:
        combined.append({
            "status": "NEW",
            "group": n["group"],
            "name": n["name"],
            "type": "array",
            "desc": n["description"],
            "enum": "",
            "items": format_proposal_item_fields(n["fields"]),
            "agents": n["agent"],
            "id": "",
            "is_new": True,
        })
    combined.sort(key=lambda d: ((d["group"] or "").lower(), (d["name"] or "").lower()))

    for d in combined:
        ws2.append([clean_cell(d["status"]), clean_cell(d["group"]), clean_cell(d["name"]),
                    clean_cell(d["type"]), clean_cell(d["desc"]), clean_cell(d["enum"]),
                    clean_cell(d["items"]), clean_cell(d["agents"]), clean_cell(d["id"])])
    # Wrap long columns + fills
    for idx, d in enumerate(combined, start=2):
        for col in (5, 6, 7):
            ws2.cell(row=idx, column=col).alignment = WRAP_TOP
        for col in (1, 2, 3, 4, 8, 9):
            ws2.cell(row=idx, column=col).alignment = Alignment(vertical="top", wrap_text=True)
        # Row height heuristic based on item-field lines and long definitions
        lines = max(1, str(ws2.cell(row=idx, column=7).value or "").count("\n") + 1,
                    str(ws2.cell(row=idx, column=5).value or "").count("\n") + 1)
        longest = max(len(str(ws2.cell(row=idx, column=5).value or "")),
                      len(str(ws2.cell(row=idx, column=7).value or "")))
        est = max(lines * 15, min(300, 15 + longest // 60 * 15))
        ws2.row_dimensions[idx].height = min(420, est)
        if d["status"] == "NEW":
            for col in range(1, len(headers2) + 1):
                ws2.cell(row=idx, column=col).fill = NEW_FILL
    ws2.auto_filter.ref = ws2.dimensions

    # ---------------- Dict formats ----------------
    ws3 = wb.create_sheet("Dict formats")
    headers3 = ["Field", "Type", "Required/Nullable", "Enum values", "Description"]
    ws3.append(headers3)
    style_header_row(ws3, len(headers3))
    set_widths(ws3, [22, 14, 18, 52, 100])
    for n in NEW_ATTRIBUTES:
        # block header
        r = ws3.max_row + 1
        ws3.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        cell = ws3.cell(row=r, column=1)
        cell.value = "{}  (Group: {} - Agent: {} - Type: array of objects)".format(
            n["name"], n["group"], n["agent"])
        cell.font = Font(bold=True)
        cell.fill = BLOCK_FILL
        cell.alignment = Alignment(vertical="center")
        for f in n["fields"]:
            ws3.append([
                clean_cell(f["name"]),
                clean_cell(f["type"]),
                clean_cell("required" if f["required"] else "nullable"),
                clean_cell(format_enum_values(f.get("enum", []))),
                clean_cell(f.get("description", "")),
            ])
        # Example rows (same examples as the definition, JSON pretty-printed)
        for i, ex in enumerate(n["examples"], start=1):
            ex_row = ws3.max_row + 1
            body = 'Client: "{}"\n->\n{}'.format(
                ex["transcript"],
                json.dumps(ex["items"], indent=2, ensure_ascii=False),
            )
            ws3.append(["Example {}".format(i), "", "", "", body])
            for col in range(1, 6):
                ws3.cell(row=ex_row, column=col).alignment = WRAP_TOP
            ws3.row_dimensions[ex_row].height = 150
        # blank row between blocks
        ws3.append(["", "", "", "", ""])
    # wrap all
    for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row, max_col=5):
        for c in row:
            c.alignment = WRAP_TOP
    ws3.auto_filter.ref = "A1:E1"

    # ---------------- Old -> New mapping ----------------
    ws4 = wb.create_sheet("Old -> New mapping")
    headers4 = ["Old group", "Old name", "Old type", "Old agent(s)",
                "New attribute", "New type value"]
    ws4.append(headers4)
    style_header_row(ws4, len(headers4))
    set_widths(ws4, [20, 34, 12, 18, 20, 48])
    for r in deleted_rows:
        new_attr, new_val = mapping_for_old(r.group_name, r.name)
        ws4.append([clean_cell(r.group_name), clean_cell(r.name), clean_cell(r.type),
                    clean_cell(", ".join(attr_agents.get(r.id, []))),
                    clean_cell(new_attr), clean_cell(new_val)])
    # 2 rows for adequacy booleans (kept as-is)
    for r in attrs:
        if r.name in ADEQUACY_NAMES:
            ws4.append([clean_cell(r.group_name), clean_cell(r.name), clean_cell(r.type),
                        clean_cell(", ".join(attr_agents.get(r.id, []))),
                        "kept as-is", ""])
    for row in ws4.iter_rows(min_row=2, max_row=ws4.max_row, max_col=6):
        for c in row:
            c.alignment = WRAP_TOP
    ws4.auto_filter.ref = ws4.dimensions

    # ---------------- JSON schema ----------------
    ws5 = wb.create_sheet("JSON schema")
    headers5 = ["Attribute", "Schema"]
    ws5.append(headers5)
    style_header_row(ws5, len(headers5))
    set_widths(ws5, [20, 135])
    for n in NEW_ATTRIBUTES:
        schema = build_json_schema(n)
        ws5.append([n["name"], json.dumps(schema, indent=2, ensure_ascii=False)])
    for idx in range(2, ws5.max_row + 1):
        ws5.cell(row=idx, column=1).alignment = Alignment(vertical="top")
        ws5.cell(row=idx, column=2).alignment = WRAP_TOP
        ws5.row_dimensions[idx].height = 260
    ws5.auto_filter.ref = "A1:B1"

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)

    # ---------------- Verify ----------------
    wb2 = load_workbook(OUT_PATH)
    print("Sheets:", wb2.sheetnames)
    for name in wb2.sheetnames:
        wsx = wb2[name]
        print("{}: {} rows x {} cols (incl. header)".format(name, wsx.max_row, wsx.max_column))
    wsx = wb2["Attributes v2"]
    counts = {"NEW": 0, "UNCHANGED": 0}
    for row in wsx.iter_rows(min_row=2, values_only=True):
        s = str(row[0] or "")
        if s == "NEW":
            counts["NEW"] += 1
        elif s == "UNCHANGED":
            counts["UNCHANGED"] += 1
    print("Status counts: NEW={} UNCHANGED={}".format(counts["NEW"], counts["UNCHANGED"]))
    print("DB now={} deleted={} after={}".format(total_now, n_deleted, n_after))
    for attr_name in ("assets", "goals"):
        for row in wsx.iter_rows(min_row=2, values_only=True):
            if (row[2] or "") == attr_name and (row[0] or "") == "NEW":
                print("=" * 80)
                print("DEFINITION [{}]:".format(attr_name))
                print(row[4])
                break
    print("Wrote {}".format(OUT_PATH))
    print("Summary: v2 (rev 2) proposal with {} NEW list attributes replacing "
          "{} attributes across 7 groups; {} kept "
          "(incl. 2 kept adequacy booleans); catalogue after = {}. Tax reverted, "
          "no outer wrapper on NEW lists.".format(
              len(NEW_ATTRIBUTES), n_deleted, total_now - n_deleted, n_after))


if __name__ == "__main__":
    main()
